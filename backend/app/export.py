import csv, io, json
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from .models import FIELDS


def safe(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(
        ("=", "+", "-", "@", "\t", "\r")
    ):
        return "'" + value
    return value


def export_rows(rows, spec):
    rows = [
        r
        for r in rows
        if (spec.ids is None or r["id"] in spec.ids)
        and all(r.get(k) is not None and r.get(k) != [] for k in spec.required)
    ]
    data = [{k: r.get(k) for k in spec.columns} for r in rows]
    if spec.format == "json":
        payload = {
            "rows": data,
            "columns": [
                {"field": k, "label": spec.headers.get(k, FIELDS[k])}
                for k in spec.columns
            ],
        }
        if spec.metadata:
            payload["field_status"] = [r.get("field_status", {}) for r in rows]
        return (
            json.dumps(payload, ensure_ascii=False, indent=2).encode(),
            "application/json",
        )
    headers = [safe(spec.headers.get(k, FIELDS[k])) for k in spec.columns]
    empty = {"blank": "", "dash": "—", "null": "null"}[spec.empty]
    values = [
        [safe(r[k]) if r[k] is not None else empty for k in spec.columns] for r in data
    ]
    if spec.format == "csv":
        out = io.StringIO()
        writer = csv.writer(out, delimiter=spec.delimiter)
        if spec.include_headers:
            writer.writerow(headers)
        writer.writerows(values)
        return (
            out.getvalue().encode("utf-8-sig" if spec.bom else "utf-8"),
            "text/csv; charset=utf-8",
        )
    wb = Workbook()
    ws = wb.active
    ws.title = "Организации"
    if spec.include_headers:
        ws.append(headers)
        ws.freeze_panes = "A2"
    for row in values:
        ws.append(row)
    if spec.include_headers:
        ws.auto_filter.ref = ws.dimensions
    for i in range(1, len(spec.columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = spec.width
    out = io.BytesIO()
    wb.save(out)
    return (
        out.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
