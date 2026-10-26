import io, json, csv
import pytest
from openpyxl import load_workbook
from backend.app.models import (
    Organization,
    FilterGroup,
    SearchSpec,
    ExportSpec,
    Geography,
)
from backend.app.filters import evaluate, matches
from backend.app.normalization import normalize
from backend.app.deduplication import dedup_key
from backend.app.export import export_rows


def org(**kw):
    return Organization(id="1", branch_id="1", name="Пример", **kw)


def test_unknown_contacts_not_absent():
    o = normalize({"id": "1", "name": "A"}, "Москва")
    rule = FilterGroup(conditions=[{"field": "telegram", "op": "absent"}])
    assert o.field_status["telegram"] == "unknown"
    assert evaluate(rule, o) is None
    assert not matches(rule, o)
    assert matches(rule, o, "include")


def test_partial_contact_groups_not_proof_of_absence():
    o = normalize({"id": "1", "name": "A", "contact_groups": []}, "Москва")
    assert o.field_status["phone"] == "unknown"


def test_contact_types_and_score():
    o = normalize(
        {
            "id": "2",
            "name": "B",
            "contact_groups": [
                {
                    "contacts": [
                        {"type": "phone", "value": "+7 000"},
                        {"type": "website", "value": "https://t.me/test"},
                        {"type": "website", "value": "https://wa.me/000"},
                    ]
                }
            ],
        },
        "Казань",
    )
    assert o.telegram == "https://t.me/test" and o.whatsapp == "https://wa.me/000"
    assert o.website is None and o.contact_score == 60


@pytest.mark.parametrize("logic,expected", [("AND", False), ("OR", True)])
def test_contact_logic(logic, expected):
    o = org(phone="123", field_status={"phone": "present", "email": "absent"})
    g = FilterGroup(
        logic=logic,
        conditions=[
            {"field": "phone", "op": "present"},
            {"field": "email", "op": "present"},
        ],
    )
    assert matches(g, o) == expected


def test_nested_unknown_logic():
    g = FilterGroup(
        conditions=[
            {"field": "phone", "op": "present"},
            {
                "logic": "OR",
                "conditions": [
                    {"field": "telegram", "op": "present"},
                    {"field": "whatsapp", "op": "present"},
                ],
            },
        ]
    )
    assert evaluate(g, org(phone="123", whatsapp="https://wa.me/123")) is True


def test_review_ranges():
    g = FilterGroup(
        conditions=[
            {"field": "reviews", "op": "gte", "value": 5},
            {"field": "rating", "op": "lte", "value": 4.5},
        ]
    )
    assert matches(g, org(reviews=10, rating=4.1))
    assert not matches(g, org(reviews=2, rating=4.1))


def test_domain_filter():
    g = FilterGroup(
        conditions=[{"field": "website", "op": "domain", "value": "example.org"}]
    )
    assert matches(g, org(website="https://www.example.org/path"))
    assert not matches(g, org(website="https://example.org.evil.test"))


def test_dedup_keeps_branches_and_missing_keys():
    a = org(org_id="chain", phone="+7 (000) 12", website="https://www.example.org/a")
    b = Organization(id="2", branch_id="2", org_id="chain", name="Пример")
    assert dedup_key(a, "branch_id") != dedup_key(b, "branch_id")
    assert dedup_key(a, "org_id") == dedup_key(b, "org_id")
    assert dedup_key(a, "phone") == "700012"
    assert dedup_key(a, "website") == "example.org"
    assert dedup_key(b, "phone") == "2"


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
def test_export_order_null_cyrillic_and_injection(fmt):
    rows = [org(phone="=1+1").model_dump()]
    spec = ExportSpec(
        format=fmt, columns=["phone", "name", "email"], headers={"phone": "Связь"}
    )
    data, mime = export_rows(rows, spec)
    if fmt == "csv":
        values = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))
        assert values[0] == ["Связь", "Название", "Email"]
        assert values[1] == ["'=1+1", "Пример", ""]
    elif fmt == "xlsx":
        ws = load_workbook(io.BytesIO(data)).active
        assert ws["A2"].value == "'=1+1"
        assert ws["B2"].value == "Пример"
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:C2"
    else:
        body = json.loads(data)
        assert list(body["rows"][0]) == ["phone", "name", "email"]
        assert body["rows"][0]["email"] is None
        assert body["field_status"] == [{}]


def test_selected_export_and_required():
    rows = [
        org(phone="123").model_dump(),
        Organization(id="2", branch_id="2", name="Б").model_dump(),
    ]
    data, _ = export_rows(
        rows, ExportSpec(format="json", columns=["name"], ids=["2"], required=["phone"])
    )
    assert json.loads(data)["rows"] == []


def test_invalid_geography_and_filters():
    with pytest.raises(ValueError):
        SearchSpec(cities=[Geography(name="Москва")], categories=["Кафе"], radius=500)
    with pytest.raises(ValueError):
        FilterGroup(conditions=[{"field": "fake", "op": "present"}])
    with pytest.raises(ValueError):
        FilterGroup(conditions=[{"field": "rating", "op": "gte", "value": "x"}])
