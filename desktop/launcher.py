"""Double-click launcher. The packaged browser UI is served only on loopback."""

from __future__ import annotations
import argparse
import contextlib
import io
import json
import logging
from logging.handlers import RotatingFileHandler
import os
import queue
from pathlib import Path
import socket
import sys
import tempfile
import threading
import time
import urllib.request
import webbrowser

APP_NAME = "2GIS Business Parser"


def resource_root() -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))


def profile_directory() -> Path:
    if os.name == "nt":
        return (
            Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
            / "2GISBusinessParser"
        )
    return (
        Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local" / "share"))
        / "2GISBusinessParser"
    )


class InstanceLock:
    """Lock an existing inode; never unlink it while another process may own it."""

    def __init__(self, directory: Path):
        self.directory = directory
        self.file = None

    def acquire(self):
        self.directory.mkdir(parents=True, exist_ok=True)
        self.file = (self.directory / "instance.lock").open("a+b")
        self.file.seek(0, 2)
        if self.file.tell() == 0:
            self.file.write(b"0")
            self.file.flush()
        self.file.seek(0)
        try:
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.file.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            return True
        except OSError:
            self.file.close()
            self.file = None
            return False

    def close(self):
        if self.file:
            self.file.seek(0)
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(self.file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(self.file.fileno(), fcntl.LOCK_UN)
            self.file.close()
            self.file = None


def configure(directory: Path, smoke=False):
    directory.mkdir(parents=True, exist_ok=True)
    os.chdir(directory)
    os.environ["DATABASE_URL"] = (
        "sqlite:///" + (directory / "data" / "parser.db").as_posix()
    )
    os.environ["DESKTOP_MODE"] = "true"
    os.environ["DESKTOP_DATA_DIR"] = str(directory)
    os.environ["ALLOWED_HOSTS"] = "127.0.0.1,localhost"
    if smoke:
        os.environ["DGIS_API_KEY"] = ""
    elif os.name == "nt" and not os.environ.get("DGIS_API_KEY"):
        from desktop.credentials import load_api_key

        try:
            os.environ["DGIS_API_KEY"] = load_api_key(directory)
        except OSError:
            pass
    handler = RotatingFileHandler(
        directory / "application.log",
        maxBytes=1024 * 1024,
        backupCount=1,
        encoding="utf-8",
    )
    handler.setFormatter(
        logging.Formatter("%(asctime)s %(levelname)s %(name)s %(message)s")
    )
    logging.basicConfig(level=logging.INFO, handlers=[handler], force=True)
    # Avoid HTTPX's URL logging; upstream URLs can contain an API key.
    logging.getLogger("httpx").setLevel(logging.WARNING)
    logging.getLogger("httpcore").setLevel(logging.WARNING)


class LocalServer:
    def __init__(self):
        self.socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        if os.name == "nt" and hasattr(socket, "SO_EXCLUSIVEADDRUSE"):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        self.socket.bind(("127.0.0.1", 0))
        self.socket.listen(128)
        self.url = f"http://127.0.0.1:{self.socket.getsockname()[1]}"
        import uvicorn
        from backend.app.api import app

        self.server = uvicorn.Server(
            uvicorn.Config(
                app,
                host="127.0.0.1",
                port=self.socket.getsockname()[1],
                loop="asyncio",
                http="h11",
                ws="none",
                lifespan="on",
                log_config=None,
                access_log=False,
            )
        )
        self.error = None
        self.thread = threading.Thread(
            target=self._run, name="parser-server", daemon=True
        )

    def _run(self):
        try:
            self.server.run(sockets=[self.socket])
        except BaseException as exc:
            self.error = exc
            logging.exception("Local server stopped")

    def start(self):
        self.thread.start()

    def wait_ready(self, timeout=30):
        deadline = time.monotonic() + timeout
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        while time.monotonic() < deadline:
            if self.error or not self.thread.is_alive():
                raise RuntimeError("Не удалось запустить сервер приложения")
            try:
                with opener.open(self.url + "/api/health", timeout=0.5) as response:
                    if json.load(response).get("status") == "ok":
                        return
            except (OSError, ValueError):
                pass
            time.sleep(0.1)
        raise RuntimeError("Приложение не успело запуститься")

    def stop(self):
        self.server.should_exit = True
        if self.thread.is_alive():
            self.thread.join(timeout=10)
        self.socket.close()


def existing_window(directory: Path):
    for _ in range(100):
        try:
            state = json.loads((directory / "runtime.json").read_text(encoding="utf-8"))
            port = int(state["port"])
            if 1 <= port <= 65535:
                opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
                with opener.open(
                    f"http://127.0.0.1:{port}/api/health", timeout=0.5
                ) as response:
                    if json.load(response).get("status") == "ok":
                        webbrowser.open(f"http://127.0.0.1:{port}")
                        return
        except (OSError, ValueError, KeyError):
            pass
        time.sleep(0.1)
    raise RuntimeError(
        "Приложение уже запускается. Повторите попытку через несколько секунд."
    )


def graphical_main(directory: Path):
    import tkinter as tk
    from tkinter import messagebox
    from tkinter import ttk

    lock = InstanceLock(directory)
    if not lock.acquire():
        existing_window(directory)
        return 0
    root = tk.Tk()
    root.title(APP_NAME)
    root.geometry("420x220")
    root.resizable(False, False)
    root.configure(bg="#ffffff")
    root.option_add("*Font", ("Segoe UI", 10))
    server = None
    tray = None
    actions = queue.Queue()
    try:
        configure(directory)
        tk.Label(
            root,
            text=APP_NAME,
            font=("Segoe UI", 15, "bold"),
            bg="#ffffff",
            fg="#111111",
        ).pack(pady=(24, 8))
        state = tk.StringVar(value="Запускаем приложение…")
        tk.Label(
            root, textvariable=state, bg="#ffffff", fg="#666666", wraplength=380
        ).pack(pady=6)
        buttons = tk.Frame(root, bg="#ffffff")
        buttons.pack(pady=17)
        open_button = ttk.Button(buttons, text="Открыть приложение", state="disabled")
        open_button.pack(side="left", padx=5)
        close_button = ttk.Button(
            buttons, text="Завершить работу", command=root.destroy
        )
        close_button.pack(side="left", padx=5)
        tk.Label(
            root,
            text="Работает в вашем браузере · данные на этом компьютере",
            font=("Segoe UI", 8),
            bg="#ffffff",
            fg="#888888",
        ).pack()
        from desktop.tray import TrayIcon

        tray = TrayIcon(actions.put)
        tray_available = tray.start()
        if tray_available:
            root.protocol("WM_DELETE_WINDOW", root.withdraw)
            ttk.Button(buttons, text="Свернуть в трей", command=root.withdraw).pack(
                side="left", padx=5
            )

        def control(action):
            try:
                opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
                headers = {}
                password = os.environ.get("APP_PASSWORD", "")
                if password:
                    import base64

                    headers["Authorization"] = (
                        "Basic "
                        + base64.b64encode(("desktop:" + password).encode()).decode()
                    )
                with opener.open(
                    urllib.request.Request(server.url + "/api/jobs", headers=headers),
                    timeout=10,
                ) as response:
                    jobs = json.load(response)
                if action == "resume" and any(
                    j["status"] in ("queued", "running") for j in jobs
                ):
                    return
                statuses = (
                    ("running", "queued")
                    if action == "pause"
                    else ("paused", "stopped")
                )
                target = next((j for j in jobs if j["status"] in statuses), None)
                if target:
                    request = urllib.request.Request(
                        server.url + "/api/jobs/" + target["id"] + "/control/" + action,
                        data=b"{}",
                        headers={**headers, "Content-Type": "application/json"},
                        method="POST",
                    )
                    with opener.open(request, timeout=10):
                        pass
            except Exception:
                logging.warning("Tray control request failed")

        def poll_actions():
            try:
                while True:
                    action = actions.get_nowait()
                    if action == "exit":
                        root.destroy()
                        return
                    if action == "window":
                        root.deiconify()
                        root.lift()
                    if action == "2gis":

                        def open_source():
                            from backend.app.browser.driver import browser_service

                            try:
                                browser_service.open_home()
                            except Exception:
                                webbrowser.open("https://2gis.ru")

                        threading.Thread(target=open_source, daemon=True).start()
                    if action == "open" and server and server.server.started:
                        webbrowser.open(server.url)
                    if (
                        action in ("pause", "resume")
                        and server
                        and server.server.started
                    ):
                        threading.Thread(
                            target=control, args=(action,), daemon=True
                        ).start()
            except queue.Empty:
                pass
            root.after(150, poll_actions)

        root.after(150, poll_actions)
        root.update()
        if not (resource_root() / "frontend" / "dist" / "index.html").exists():
            raise RuntimeError("В сборке отсутствует интерфейс приложения")
        server = LocalServer()
        server.start()

        def ready():
            if server.error or not server.thread.is_alive():
                state.set("Ошибка запуска. Подробности в application.log.")
                return
            if server.server.started:
                (directory / "runtime.json").write_text(
                    json.dumps({"port": server.socket.getsockname()[1]}),
                    encoding="utf-8",
                )
                state.set(
                    "Приложение работает. Для остановки нажмите «Завершить работу»."
                )
                open_button.configure(
                    state="normal", command=lambda: webbrowser.open(server.url)
                )
                webbrowser.open(server.url)
                if tray_available:
                    root.after(500, root.withdraw)
            else:
                root.after(100, ready)

        root.after(100, ready)
        root.mainloop()
        return 0
    except Exception:
        logging.exception("Desktop launch failed")
        messagebox.showerror(
            APP_NAME,
            "Не удалось запустить приложение.\nПодробности: "
            + str(directory / "application.log"),
        )
        return 1
    finally:
        if tray:
            tray.stop()
        if server:
            server.stop()
        (directory / "runtime.json").unlink(missing_ok=True)
        lock.close()
        with contextlib.suppress(tk.TclError):
            root.destroy()


def smoke_test(result_path: Path):
    """Exercise the packaged server and XLSX export, using only a fresh temporary DB."""
    old_cwd = Path.cwd()
    report = {"ok": False}
    try:
        with tempfile.TemporaryDirectory(prefix="parser-exe-test-") as temporary:
            directory = Path(temporary)
            configure(directory, smoke=True)
            server = LocalServer()
            try:
                server.start()
                server.wait_ready()
                import httpx
                from openpyxl import load_workbook

                with httpx.Client(base_url=server.url, trust_env=False) as client:
                    response = client.get("/")
                    response.raise_for_status()
                    assert '<html lang="ru">' in response.text
                    response = client.post("/api/connection", json={"mode": "demo"})
                    response.raise_for_status()
                    response = client.post(
                        "/api/jobs",
                        json={
                            "mode": "demo",
                            "cities": [{"name": "Москва"}],
                            "categories": ["Кафе"],
                            "limit": 10,
                            "requests_per_second": 5,
                        },
                    )
                    response.raise_for_status()
                    job_id = response.json()["id"]
                    for _ in range(100):
                        job = client.get("/api/jobs/" + job_id).json()
                        if job["status"] == "completed":
                            break
                        if job["status"] == "failed":
                            raise RuntimeError(
                                "Демонстрационная задача завершилась с ошибкой"
                            )
                        time.sleep(0.1)
                    assert job["status"] == "completed" and job["matched"] == 10
                    response = client.post(
                        "/api/jobs/" + job_id + "/export",
                        json={"format": "xlsx", "columns": ["name", "phone"]},
                    )
                    response.raise_for_status()
                    sheet = load_workbook(io.BytesIO(response.content)).active
                    assert sheet.max_row == 11 and sheet.max_column == 2
                    credential_test = "not-applicable"
                    if os.name == "nt":
                        from desktop.credentials import save_api_key, load_api_key

                        save_api_key("local-test-value", directory)
                        assert load_api_key(directory) == "local-test-value"
                        assert (
                            b"local-test-value"
                            not in (directory / "api-key.dpapi").read_bytes()
                        )
                        credential_test = "passed"
                    tray_test = "not-applicable"
                    if os.name == "nt":
                        from desktop.tray import TrayIcon

                        tray = TrayIcon(lambda action: None)
                        try:
                            tray.start()
                            assert tray.window_created
                            tray_test = "passed"
                        finally:
                            tray.stop()
                        assert not tray.thread.is_alive()
                    browser_test = "not-requested"
                    if os.getenv("DGIS_BROWSER_TEST_PORT"):
                        import asyncio
                        from backend.app.browser.driver import BrowserService
                        from backend.app.providers.browser import BrowserProvider
                        from backend.app.models import SearchSpec

                        service = BrowserService()
                        service.port = int(os.environ["DGIS_BROWSER_TEST_PORT"])
                        provider = BrowserProvider(service)
                        provider.wait = lambda seconds: time.sleep(min(seconds, 0.15))
                        specification = SearchSpec(
                            mode="browser",
                            cities=[{"name": "Москва", "region_id": "moscow"}],
                            categories=["Кафе"],
                        )

                        async def browser_check():
                            rows = []
                            for page in range(1, 6):
                                found, more = await provider.page(
                                    specification, specification.cities[0], "Кафе", page
                                )
                                rows += found
                                if not more:
                                    break
                            assert [o.id for o in rows] == ["101", "102", "103"]
                            assert all(
                                o.phone == "+70000000000"
                                and o.website.rstrip("/") == "https://example.org"
                                and o.telegram == "https://t.me/example"
                                for o in rows
                            )

                        try:
                            asyncio.run(browser_check())
                            browser_test = "passed"
                        finally:
                            service.close()
                    report = {
                        "ok": True,
                        "matched": 10,
                        "xlsx_rows": 10,
                        "xlsx_columns": 2,
                        "credential_storage": credential_test,
                        "tray_window": tray_test,
                        "browser_cdp": browser_test,
                        "frozen": bool(getattr(sys, "frozen", False)),
                    }
            finally:
                server.stop()
                from backend.app.database import engine

                engine.dispose()
                logging.shutdown()
                os.chdir(old_cwd)
    except Exception as exc:
        report = {"ok": False, "error_type": type(exc).__name__}
    result_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return 0 if report["ok"] else 1


def main():
    parser = argparse.ArgumentParser(description=APP_NAME)
    parser.add_argument("--smoke-test", action="store_true")
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    if args.smoke_test:
        if args.result is None:
            parser.error("--smoke-test requires --result")
        return smoke_test(args.result.resolve())
    try:
        return graphical_main(profile_directory())
    except Exception as error:
        if os.name == "nt":
            import ctypes

            ctypes.windll.user32.MessageBoxW(None, str(error), APP_NAME, 0x10)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
